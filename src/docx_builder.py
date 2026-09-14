"""
Stage 4b: renders a GeneratedDocument to an actual .docx file, following the
formatting table in Section 5 of the format-explainer document:
  - forum heading / jurisdiction / case number / affidavit title: bold, caps, centred
  - cause title party names: normal, left; status tags: right-aligned; VERSUS: centred
  - paragraph numbers: bold; paragraph text: normal, justified
  - PRAYER / VERIFICATION headings: bold, caps, centred; prayer letters: bold
  - DEPONENT: caps, right-aligned; Before Me: left-aligned
"""
from __future__ import annotations
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from schemas import GeneratedDocument
from date_utils import to_ordinal_date


def _add_centered_bold_caps(doc: Document, text: str, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(size)
    return p


def _add_left_right_pair(doc: Document, left_text: str, right_text: str):
    """Party name on the left, status tag right-aligned -- approximated
    with a two-cell borderless table so both sit on the same line."""
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    left_cell, right_cell = table.rows[0].cells
    left_cell.paragraphs[0].add_run(left_text)
    rp = right_cell.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.add_run(right_text)
    return table


def build_docx(
    generated: GeneratedDocument,
    verification_range_end: int,
    output_path: str,
):
    e = generated.entities
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    # 1. Forum heading
    _add_centered_bold_caps(doc, e.forum)
    # 2. Jurisdiction
    _add_centered_bold_caps(doc, e.jurisdiction_type)
    # 3. Case number
    _add_centered_bold_caps(doc, f"{e.proceeding_type} NO. {e.case_number} OF {e.year}")

    doc.add_paragraph()

    # 4. Cause title
    _add_left_right_pair(doc, e.petitioner, "...Petitioner")
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("VERSUS").bold = True
    doc.add_paragraph()
    for r in e.respondents:
        _add_left_right_pair(doc, f"{r.number}. {r.name}", r.status_tag)
    doc.add_paragraph()

    # 5. Affidavit title
    _add_centered_bold_caps(doc, f"AFFIDAVIT IN REPLY ON BEHALF OF RESPONDENT NO. {e.respondent_number}")
    doc.add_paragraph()

    # 6. Deponent clause
    if e.is_organisation:
        role = f"the {e.designation} of the Respondent No.{e.respondent_number} above named"
    else:
        role = f"the Respondent No.{e.respondent_number} above named"
    clause = (
        f"I, {e.deponent_name}, {(e.designation + ', ') if e.is_organisation else ''}"
        f"residing at {e.address}, {role}, do hereby {e.verification_verb} and state as under:"
    )
    doc.add_paragraph(clause)
    doc.add_paragraph()

    # 7. Numbered paragraphs
    for i, para_text in enumerate(generated.paragraphs, start=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        num_run = p.add_run(f"{i}. ")
        num_run.bold = True
        p.add_run(para_text)

    doc.add_paragraph()

    # 8. Prayer
    _add_centered_bold_caps(doc, "PRAYER")
    doc.add_paragraph(
        f"I therefore respectfully pray that this Hon'ble Court may be pleased to:"
    )
    letters = "abcdefgh"
    for i, line in enumerate(generated.prayer_lines):
        p = doc.add_paragraph()
        letter_run = p.add_run(f"({letters[i]}) ")
        letter_run.bold = True
        p.add_run(line)

    doc.add_paragraph()

    # 9. Jurat
    verb_map = {"solemnly affirm": "Solemnly affirmed", "swear and affirm": "Sworn"}
    jurat_verb = verb_map.get(e.verification_verb.lower(), "Solemnly affirmed")
    ordinal_date = to_ordinal_date(e.date_of_attestation)
    doc.add_paragraph(f"{jurat_verb} at {e.place_of_attestation}")
    doc.add_paragraph(f"On this {ordinal_date}")
    table = doc.add_table(rows=1, cols=2)
    left_cell, right_cell = table.rows[0].cells
    left_cell.paragraphs[0].add_run("Before Me")
    rp = right_cell.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.add_run("DEPONENT").bold = True

    doc.add_paragraph()

    # 10. Verification
    _add_centered_bold_caps(doc, "VERIFICATION")
    verification_text = (
        f"I, {e.deponent_name}, the Deponent above named, do hereby verify that the "
        f"contents of paragraphs 1 to {verification_range_end} and the Prayer above are "
        f"true and correct to my knowledge and belief and that nothing material has been "
        f"concealed therefrom."
    )
    doc.add_paragraph(verification_text)
    doc.add_paragraph(f"Verified at {e.place_of_attestation} on this {ordinal_date}.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("DEPONENT").bold = True

    # Advocate / drafting block
    if e.advocate_firm:
        doc.add_paragraph()
        doc.add_paragraph(e.advocate_firm.upper())
        doc.add_paragraph(f"Advocates for the Respondent No.{e.respondent_number}.")

    doc.save(output_path)
    return output_path
